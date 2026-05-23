#!/usr/bin/env python3
"""PPTX ingestion: slide text + embedded images + per-slide vision captions.

Companion to ``ingest_pdf.py`` (ROADMAP §1, LOOP_QUEUE item 13). Same artifact
lifecycle (``queued → extracting → chunking → embedded → ready``), same
``source_chunks`` table, same glossary/context merge — adapted for slide
decks:

* per-slide text frames are concatenated and chunked at slide granularity
* each embedded picture is written to ``<source>/images/slide{N}_img{K}.{ext}``
  and captioned via ``_caption.caption_image``
* the per-slide caption (text + image descriptions) is what lands in
  ``source_chunks.text`` so retrieval over a deck recovers slides intact
* the final ``view_spec`` is a markdown document the universal renderer shows
  on ``/app/artifacts/<id>`` (title + slide-by-slide layout with the captions)
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import traceback
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from _caption import caption_image  # noqa: E402
from ingest_pdf import (  # noqa: E402
    EMBEDDING_DIM,
    _fetch_artifact,
    _insert_chunks,
    _load_env,
    _patch_artifact,
    _set_failed,
    _set_pipeline,
    build_diff_view_spec,
    embed_text,
    extract_context,
    extract_glossary,
    merge_context,
    merge_glossary,
)


def _extract_slides(pptx_path: Path, images_dir: Path) -> list[dict[str, Any]]:
    """Return a list of slides, each carrying its text, captions, and image paths."""
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "python-pptx not installed in the venv — run "
            "`pip install --target=/home/agent/zap/.venv/lib/python3.12/site-packages python-pptx`"
        ) from exc

    pres = Presentation(str(pptx_path))
    images_dir.mkdir(parents=True, exist_ok=True)

    slides: list[dict[str, Any]] = []
    for slide_no, slide in enumerate(pres.slides, start=1):
        texts: list[str] = []
        images: list[dict[str, Any]] = []
        for shape_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs).strip()
                    if line:
                        texts.append(line)
            elif getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE = 13
                try:
                    image = shape.image
                except Exception:
                    continue
                ext = image.ext.lower() if hasattr(image, "ext") else "png"
                img_path = images_dir / f"slide{slide_no:03d}_img{shape_idx:02d}.{ext}"
                img_path.write_bytes(image.blob)
                images.append({"path": str(img_path), "ext": ext})

        if hasattr(slide, "notes_slide") and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide else ""
            if notes:
                texts.append(f"[notes] {notes}")

        # Caption each image with the slide's title/text as context so the
        # vision model can disambiguate "this slide's chart" vs. a random photo.
        slide_context = texts[0] if texts else ""
        captions: list[str] = []
        for img in images:
            cap = caption_image(
                Path(img["path"]),
                context=f"This is slide {slide_no} of a presentation. Slide text: {slide_context!r}",
            )
            captions.append(cap)
            img["caption"] = cap

        slides.append(
            {
                "slide": slide_no,
                "text": "\n".join(texts).strip(),
                "captions": captions,
                "images": images,
            }
        )
    return slides


def _build_chunks(slides: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One chunk per slide. Combine slide text + image captions verbatim."""
    chunks: list[dict[str, Any]] = []
    for ordinal, s in enumerate(slides):
        parts: list[str] = []
        if s["text"]:
            parts.append(s["text"])
        for i, cap in enumerate(s["captions"], start=1):
            parts.append(f"[image {i}] {cap}")
        body = "\n\n".join(parts).strip()
        if not body:
            body = f"[empty slide {s['slide']}]"
        chunks.append(
            {
                "ordinal": ordinal,
                "page": s["slide"],
                "text": body,
                "metadata": {"image_count": len(s["images"])},
            }
        )
    return chunks


def _insert_chunks_with_meta(env, artifact, chunks):
    """Mirror of ingest_pdf._insert_chunks but preserves per-chunk metadata."""
    import urllib.error
    import urllib.request

    user_id = artifact.get("user_id")
    artifact_id = artifact["id"]
    try:
        from ingest_pdf import _supabase_request

        _supabase_request(env, "DELETE", f"source_chunks?artifact_id=eq.{artifact_id}")
    except urllib.error.HTTPError:
        pass

    if not chunks:
        return

    BATCH = 50
    from ingest_pdf import _supabase_request

    for i in range(0, len(chunks), BATCH):
        rows = [
            {
                "artifact_id": artifact_id,
                "user_id": user_id,
                "ordinal": c["ordinal"],
                "page": c.get("page"),
                "text": c["text"],
                "embedding": embed_text(c["text"]),
                "metadata": c.get("metadata", {}),
            }
            for c in chunks[i : i + BATCH]
        ]
        _supabase_request(env, "POST", "source_chunks", rows)


def _summary_markdown(artifact: dict[str, Any], slides: list[dict[str, Any]]) -> str:
    name = artifact.get("name") or artifact.get("slug") or "Slide deck"
    bits: list[str] = [f"# {name}", f"_{len(slides)} slides ingested._"]
    for s in slides[:50]:  # cap for the inline summary view
        bits.append(f"## Slide {s['slide']}")
        if s["text"]:
            bits.append(s["text"])
        for i, cap in enumerate(s["captions"], start=1):
            bits.append(f"**Image {i}.** {cap}")
    if len(slides) > 50:
        bits.append(f"\n_{len(slides) - 50} more slides — see the source chunks table._")
    return "\n\n".join(bits)


def ingest_pptx(artifact_id: str, pptx_path: Path, workspace: Path) -> dict[str, Any]:
    env = _load_env()
    if "SUPABASE_SERVICE_ROLE_KEY" not in env:
        raise RuntimeError("SUPABASE_SERVICE_ROLE_KEY missing (env or .env.local)")

    artifact = _fetch_artifact(env, artifact_id)
    _set_pipeline(env, artifact, "extracting")

    images_dir = pptx_path.parent / "images"
    slides = _extract_slides(pptx_path, images_dir)
    if not slides:
        raise RuntimeError("no slides in PPTX")

    _set_pipeline(env, artifact, "chunking")
    chunks = _build_chunks(slides)

    full_text = "\n\n".join(c["text"] for c in chunks)

    _set_pipeline(env, artifact, "embedded")
    _insert_chunks_with_meta(env, artifact, chunks)

    terms = extract_glossary(full_text)
    blurb = extract_context(full_text)
    source_label = artifact.get("name") or artifact.get("slug") or artifact_id

    merge_glossary(workspace / "glossary.md", terms, source_label)
    merge_context(workspace / "company-context.md", blurb, source_label)

    diff_view = build_diff_view_spec(env, artifact, full_text) or {}

    meta = dict(artifact.get("metadata") or {})
    meta.update(
        {
            "pipeline_status": "ready",
            "slides": len(slides),
            "images": sum(len(s["images"]) for s in slides),
            "characters": len(full_text),
            "chunks": len(chunks),
            "extracted_at": datetime.now(timezone.utc).isoformat(),
            "extracted_text": full_text[:200_000],
            "glossary_terms_added": len(terms),
            "embedding_dim": EMBEDDING_DIM,
            "embedding_kind": "hash-bow-v1",
        }
    )
    view_spec = {
        "renderer": "markdown",
        "text": _summary_markdown(artifact, slides),
    }
    _patch_artifact(
        env,
        artifact_id,
        {
            "status": "draft",
            "metadata": meta,
            "view_spec": {**(artifact.get("view_spec") or {}), **view_spec, **diff_view},
        },
    )
    return {
        "artifact_id": artifact_id,
        "slides": len(slides),
        "images": meta["images"],
        "chunks": len(chunks),
        "terms": len(terms),
    }


def _cli() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--artifact-id", required=True)
    p.add_argument("--pptx", required=True)
    p.add_argument("--workspace", required=True)
    args = p.parse_args()
    try:
        result = ingest_pptx(args.artifact_id, Path(args.pptx), Path(args.workspace))
    except Exception as exc:
        try:
            env = _load_env()
            artifact = _fetch_artifact(env, args.artifact_id)
            _set_failed(env, artifact, f"{type(exc).__name__}: {exc}\n{traceback.format_exc()}")
        except Exception:
            pass
        sys.stderr.write(f"ingest_pptx failed: {exc}\n{traceback.format_exc()}")
        return 1
    sys.stdout.write(json.dumps(result) + "\n")
    return 0


# Silence unused-import warning while keeping the import for documentation.
_ = re

if __name__ == "__main__":
    sys.exit(_cli())
